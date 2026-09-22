from __future__ import annotations

from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any
import re


class DocumentRenderer(ABC):
    format: str

    @abstractmethod
    def render(self, specification: dict[str, Any], output_path: str) -> str:
        """Render validated structured content deterministically."""
        raise NotImplementedError


class PDFRenderer(DocumentRenderer):
    format = "pdf"

    def render(self, specification: dict[str, Any], output_path: str) -> str:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
        from xml.sax.saxutils import escape

        text = str(
            specification.get("text")
            or specification.get("content")
            or ""
        ).strip()

        if not text:
            raise ValueError("PDF rendering requires document text")

        output = Path(output_path).expanduser()
        output.parent.mkdir(parents=True, exist_ok=True)

        document = SimpleDocTemplate(
            str(output),
            pagesize=A4,
            rightMargin=50,
            leftMargin=50,
            topMargin=50,
            bottomMargin=50,
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "SwicoTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            spaceAfter=18,
        )

        heading_style = ParagraphStyle(
            "SwicoHeading",
            parent=styles["Heading2"],
            spaceBefore=12,
            spaceAfter=6,
        )

        body_style = ParagraphStyle(
            "SwicoBody",
            parent=styles["BodyText"],
            leading=15,
            spaceAfter=7,
        )

        story = []
        title_added = False

        heading_pattern = re.compile(
            r"^(introduction|objectives|system architecture|"
            r"implementation approach|implementation|conclusion|"
            r"summary|overview|background|methodology|results|discussion|"
            r"references)\s*:?[ ]*$",
            re.IGNORECASE,
        )

        for raw_line in text.splitlines():
            line = raw_line.strip()

            if not line:
                story.append(Spacer(1, 6))
                continue

            if not title_added and re.match(
                r"^(title|document title)\s*:",
                line,
                re.IGNORECASE,
            ):
                title = line.split(":", 1)[1].strip()
                story.append(Paragraph(escape(title), title_style))
                title_added = True
                continue

            if heading_pattern.match(line):
                story.append(
                    Paragraph(
                        escape(line.rstrip(":")),
                        heading_style,
                    )
                )
                continue

            if re.match(r"^[-*•]\s+", line):
                item = re.sub(r"^[-*•]\s+", "", line)
                story.append(
                    ListFlowable(
                        [ListItem(Paragraph(escape(item), body_style))],
                        bulletType="bullet",
                        leftIndent=18,
                    )
                )
                continue

            if re.match(r"^\d+[.)]\s+", line):
                item = re.sub(r"^\d+[.)]\s+", "", line)
                story.append(
                    ListFlowable(
                        [ListItem(Paragraph(escape(item), body_style))],
                        bulletType="1",
                        start="1",
                        leftIndent=18,
                    )
                )
                continue

            story.append(Paragraph(escape(line), body_style))

        document.build(story)
        return str(output)


class DOCXRenderer(DocumentRenderer):
    format = "docx"

    def render(self, specification: dict[str, Any], output_path: str) -> str:
        from docx import Document

        text = str(
            specification.get("text")
            or specification.get("content")
            or ""
        ).strip()

        if not text:
            raise ValueError("DOCX rendering requires document text")

        output = Path(output_path).expanduser()
        output.parent.mkdir(parents=True, exist_ok=True)

        document = Document()

        lines = text.splitlines()
        title_added = False

        for raw_line in lines:
            line = raw_line.strip()

            if not line:
                continue

            if not title_added and re.match(
                r"^(title|document title)\s*:",
                line,
                re.IGNORECASE,
            ):
                title = line.split(":", 1)[1].strip()
                document.add_heading(title, level=0)
                title_added = True
                continue

            if re.match(r"^(introduction|objectives|system architecture|"
                        r"conclusion|summary|overview|background|"
                        r"methodology|results|discussion|references)\s*$",
                        line,
                        re.IGNORECASE):
                document.add_heading(line, level=1)
                continue

            if re.match(r"^\d+\.\s+", line):
                document.add_paragraph(
                    re.sub(r"^\d+\.\s+", "", line),
                    style="List Number",
                )
                continue

            if re.match(r"^[-*]\s+", line):
                document.add_paragraph(
                    re.sub(r"^[-*]\s+", "", line),
                    style="List Bullet",
                )
                continue

            document.add_paragraph(line)

        document.save(output)

        return str(output)


class PPTXRenderer(DocumentRenderer):
    format = "pptx"

    def render(self, specification: dict[str, Any], output_path: str) -> str:
        from pptx import Presentation

        text = str(
            specification.get("text")
            or specification.get("content")
            or ""
        ).strip()

        if not text:
            raise ValueError("PPTX rendering requires document text")

        output = Path(output_path).expanduser()
        output.parent.mkdir(parents=True, exist_ok=True)

        presentation = Presentation()

        # Remove the default empty slide if possible.
        if presentation.slides:
            slide_id = presentation.slides._sldIdLst[0].rId
            presentation.part.drop_rel(slide_id)
            del presentation.slides._sldIdLst[0]

        lines = [line.strip() for line in text.splitlines() if line.strip()]

        title = "AI Assistant Project"
        body_lines = lines

        if lines and re.match(
            r"^(title|document title)\s*:",
            lines[0],
            re.IGNORECASE,
        ):
            title = lines[0].split(":", 1)[1].strip()
            body_lines = lines[1:]

        # Title slide
        slide = presentation.slides.add_slide(
            presentation.slide_layouts[0]
        )
        slide.shapes.title.text = title
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = "Generated by Swico + SmolLM2-360M"

        # Content slides
        current_title = "Overview"
        current_body = []

        def add_content_slide(slide_title: str, content: list[str]) -> None:
            if not content:
                return

            slide = presentation.slides.add_slide(
                presentation.slide_layouts[1]
            )
            slide.shapes.title.text = slide_title

            frame = slide.placeholders[1].text_frame
            frame.clear()

            for index, item in enumerate(content):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = item
                paragraph.level = 0

        heading_pattern = re.compile(
            r"^(introduction|objectives|system architecture|"
            r"implementation approach|implementation|conclusion|"
            r"summary|overview|background|methodology|results|discussion|"
            r"references)\s*:?[ ]*$",
            re.IGNORECASE,
        )

        for line in body_lines:
            if heading_pattern.match(line):
                add_content_slide(current_title, current_body)
                current_title = line.rstrip(":").strip()
                current_body = []
            else:
                clean = re.sub(r"^[-*•]\s+", "", line)
                clean = re.sub(r"^\d+[.)]\s+", "", clean)
                current_body.append(clean)

                if len(current_body) >= 6:
                    add_content_slide(current_title, current_body)
                    current_body = []

        add_content_slide(current_title, current_body)

        presentation.save(output)
        return str(output)


class XLSXRenderer(DocumentRenderer):
    format = "xlsx"
